import re, sys, os
from docx import Document
from docx.shared import Inches, Pt
for md in sys.argv[1:]:
    base = os.path.dirname(md)
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    lines = open(md, encoding="utf8").read().splitlines()
    doc.add_heading(os.path.basename(base).replace("(DONE) ", ""), level=1)
    para = None
    for raw in lines:
        line = raw.strip()
        m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if m:
            doc.add_picture(os.path.join(base, m.group(2)), width=Inches(6.2))
            cap = doc.add_paragraph(m.group(1)); cap.runs[0].italic = True; cap.runs[0].font.size = Pt(9)
            para = None; continue
        if not line:
            para = None; continue
        if re.match(r"^(Flag|Category|Tools used|Key insight|Steps|Note):", line):
            k, _, v = line.partition(":")
            p = doc.add_paragraph(); p.add_run(k + ":").bold = True; p.add_run(v)
            para = p; continue
        if re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(); p.add_run(line).bold = True; para = None; continue
        if line.startswith("- "):
            para = doc.add_paragraph(line[2:], style="List Bullet"); continue
        if para is not None:  # continuation of a wrapped line
            para.add_run(" " + line)
        else:
            para = doc.add_paragraph(line)
    out = os.path.join(base, "writeup.docx"); doc.save(out); print("saved", out)
